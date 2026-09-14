"""
Tests for the Breakdown Export.

Two formats written from one set of rows, plus the thumbnails the rows point
at. Nothing here calls a model: an export is the descriptions written out, and
they already exist by the time anyone presses the button.
"""

import csv

from openpyxl import load_workbook

from src.core.models import Interpretation, ShotRecord
from src.identifier.export import (
    EXPORT_COLUMNS,
    EXPORT_CSV,
    EXPORT_XLSX,
    THUMBNAIL_DIR,
    export_rows,
    write_csv,
    write_thumbnails,
    write_xlsx,
)
from src.identifier.pipeline import FRAMES_DIRECTORY

DESCRIBED = Interpretation(
    shot_size="EWS",
    shot_type="single",
    characters=["Nico"],
    location="the jetty",
    camera_move="crane up",
    summary="EWS single on Nico at the controls of a small boat.",
)


def described_shot(directory, name="shot_001.mp4", number="PARA_003_4560_blockout_v0001"):
    path = directory / name
    path.write_bytes(b"not really a video")

    return ShotRecord(
        file=str(path), shot_number=number, interpretation=DESCRIBED, confidence=0.82
    )


def sampled_frame(directory, stem):
    """The frame the vision pass would have left behind."""
    frames = directory / FRAMES_DIRECTORY
    frames.mkdir(exist_ok=True)

    frame = frames / f"{stem}_f00.jpg"
    frame.write_bytes(b"jpeg-ish")
    return frame


# --- THE ROWS BOTH FORMATS ARE BUILT FROM ---


def test_a_row_carries_the_reading_in_the_export_columns(tmp_path):
    row = export_rows([described_shot(tmp_path)])[0]

    assert set(row) == set(EXPORT_COLUMNS)
    assert row["shot_number"] == "PARA_003_4560_blockout_v0001"
    assert row["characters"] == "Nico"
    assert row["description"].startswith("EWS single on Nico")


def test_a_shot_with_no_number_is_still_a_row(tmp_path):
    """
    A breakdown exists to give unnamed shots their numbers. Dropping them
    would remove exactly the rows the user needs to act on.
    """
    record = described_shot(tmp_path, number=None)

    rows = export_rows([record])

    assert len(rows) == 1
    assert rows[0]["shot_number"] == ""


def test_a_shot_that_was_never_described_is_still_a_row(tmp_path):
    record = ShotRecord(file=str(tmp_path / "unread.mp4"))

    rows = export_rows([record])

    assert rows[0]["description"] == ""
    assert rows[0]["characters"] == ""


# --- CSV ---


def test_the_csv_has_a_header_and_one_row_per_shot(tmp_path):
    records = [described_shot(tmp_path, "a.mp4"), described_shot(tmp_path, "b.mp4", None)]

    target = write_csv(records, tmp_path / "breakdown")

    assert target.name == EXPORT_CSV
    with open(target, newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))

    assert len(rows) == 2
    assert list(rows[0]) == list(EXPORT_COLUMNS)


def test_the_csv_opens_without_blank_rows_between(tmp_path):
    """
    Windows translates the csv module's own line endings again unless it is
    told not to, which puts an empty row between every real one.
    """
    target = write_csv([described_shot(tmp_path)], tmp_path / "breakdown")

    text = target.read_text(encoding="utf-8-sig")

    assert "\r\r\n" not in text


# --- SPREADSHEET ---


def test_the_spreadsheet_says_the_same_as_the_csv(tmp_path):
    """
    Both come from `export_rows`, and this is what keeps that true — the two
    files are a format apart, never a content apart.
    """
    records = [described_shot(tmp_path, "a.mp4"), described_shot(tmp_path, "b.mp4", None)]
    output = tmp_path / "breakdown"

    write_csv(records, output)
    xlsx = write_xlsx(records, output)

    with open(output / EXPORT_CSV, newline="", encoding="utf-8-sig") as handle:
        from_csv = [list(row.values()) for row in csv.DictReader(handle)]

    sheet = load_workbook(xlsx).active
    from_sheet = [[cell if cell is not None else "" for cell in row]
                  for row in sheet.iter_rows(min_row=2, values_only=True)]

    assert from_sheet == from_csv


def test_the_spreadsheet_header_is_bold_and_stays_put(tmp_path):
    """Forty rows is more than a screen, and a scrolled sheet with no header is unreadable."""
    xlsx = write_xlsx([described_shot(tmp_path)], tmp_path / "breakdown")

    sheet = load_workbook(xlsx).active

    assert xlsx.name == EXPORT_XLSX
    assert sheet["A1"].font.bold
    assert sheet.freeze_panes == "A2"
    assert sheet.title == "Shot breakdown"


# --- THUMBNAILS ---


def test_a_thumbnail_is_copied_from_the_frame_already_sampled(tmp_path):
    """The vision pass has decoded this shot once. It does not get decoded again."""
    record = described_shot(tmp_path)
    sampled_frame(tmp_path, "shot_001")

    written = write_thumbnails([record], tmp_path / "breakdown")

    assert len(written) == 1
    assert written[0].name == "PARA_003_4560_blockout_v0001.jpg"
    assert written[0].read_bytes() == b"jpeg-ish"


def test_the_thumbnail_column_points_at_the_file_that_was_written(tmp_path):
    record = described_shot(tmp_path)
    sampled_frame(tmp_path, "shot_001")
    output = tmp_path / "breakdown"

    write_thumbnails([record], output)
    row = export_rows([record])[0]

    assert (output / row["thumbnail"]).is_file()
    assert row["thumbnail"].startswith(f"{THUMBNAIL_DIR}/")


def test_a_renamed_shot_still_finds_its_frames(tmp_path):
    """
    Frames are named after the file as it was when it was sampled. Looking
    only under the current name would find nothing for exactly the shots that
    have been through the rename step.
    """
    record = described_shot(tmp_path, "PARA_003_4560_blockout_v0001.mp4")
    record.original_file = "shot_001.mp4"
    sampled_frame(tmp_path, "shot_001")

    written = write_thumbnails([record], tmp_path / "breakdown")

    assert len(written) == 1


def test_a_shot_with_no_frames_left_is_skipped_not_invented(tmp_path):
    record = described_shot(tmp_path)

    written = write_thumbnails([record], tmp_path / "breakdown")

    assert written == []
